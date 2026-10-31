"""Чтение реестра ППУ из Excel.

Лист и строку заголовка ищем сами: в реальном реестре они могут
отличаться от примера (в примере данные на «Лист2», «Лист1» пустой).
"""
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

# Каноническое имя поля -> варианты заголовка в реестре (после _norm_header)
COLUMNS = {
    "reg_no": ["рег. №", "рег.№", "рег №", "регистрационный номер", "№", "№ п/п"],
    "name": ["наименование", "наименование ппу"],
    "before": ["до улучшения"],
    "after": ["после улучшения"],
    "proposal": ["предложения", "предложение"],
    "result": ["результат"],
    "status": ["статус"],
    "effect": ["ээ", "экономический эффект"],
    "date_submitted": ["дата подачи"],
    "date_registered": ["дата регистрации"],
    "date_reviewed": ["дата рассмотрения"],
    "date_implemented": ["дата реализации"],
    "date_effect": ["дата ээ"],
}
REQUIRED = ("name", "before", "after", "proposal")
HEADER_SCAN_ROWS = 10


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _match_columns(header_row) -> dict[str, int]:
    """Канонические поля -> индекс колонки."""
    lookup = {alias: field for field, aliases in COLUMNS.items() for alias in aliases}
    mapping = {}
    for idx, cell in enumerate(header_row):
        field = lookup.get(_norm_header(cell))
        if field and field not in mapping:
            mapping[field] = idx
    return mapping


def _find_table(wb):
    """Первый лист, где в первых строках есть заголовок с нужными колонками."""
    for ws in wb.worksheets:
        for row_idx, row in enumerate(
            ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True), start=1
        ):
            mapping = _match_columns(row)
            if all(f in mapping for f in REQUIRED):
                return ws, row_idx, mapping
    raise ValueError(f"Не найден лист с колонками: {', '.join(REQUIRED)}")


def _convert(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def load_registry(path: str | Path) -> list[dict]:
    """Строки реестра как словари с каноническими полями.

    У каждой записи есть `id`: регистрационный номер, а если его нет
    (ещё не зарегистрированные ППУ) — `row<номер строки>`.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws, header_row, mapping = _find_table(wb)
        records = []
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            rec = {field: _convert(row[idx]) if idx < len(row) else None
                   for field, idx in mapping.items()}
            if not any(rec.get(f) for f in REQUIRED):
                continue
            reg_no = rec.get("reg_no")
            if isinstance(reg_no, float) and reg_no.is_integer():
                reg_no = rec["reg_no"] = int(reg_no)
            rec["id"] = str(reg_no) if reg_no is not None else f"row{row_idx}"
            rec["sheet"] = ws.title
            rec["row"] = row_idx
            records.append(rec)
        return records
    finally:
        wb.close()
